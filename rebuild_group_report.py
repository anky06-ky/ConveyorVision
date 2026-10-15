from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm, inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image, PageBreak, PageTemplate, Paragraph, SimpleDocTemplate, Spacer,
    Table, TableStyle, Frame, BaseDocTemplate, KeepTogether,
)

OUTDIR = Path(r'D:\Nhu\BT_03')
FULL_DOCX = OUTDIR / 'Bai_tap_nhom_5_Ban_hoan_chinh_10_trang_TNR13_Final.docx'
PDF_OUT = OUTDIR / 'Bai_tap_nhom_5_Ban_hoan_chinh_10_trang_TNR13_Final.pdf'
LOGO = Path(r'C:\Users\trana\AppData\Local\Temp\report_logo_full.png')

MEMBERS = [
    ['1', 'Nguyễn Quỳnh Tâm Như', '2473801070142', 'Làm Word, hoàn thành báo cáo', '100%'],
    ['2', 'Nguyễn Thùy Hoàng Duyên', '2473801070238', 'Tổng hợp nội dung', '100%'],
    ['3', 'Nguyễn Thị Như Huyền', '2373801070074', 'Nội dung', '100%'],
    ['4', 'Nguyễn Thị Thanh Huyền', '2373801071480', 'Nội dung', '100%'],
    ['5', 'Trần Hoàng Phương Quỳnh', '2173801010109', 'Nội dung', '100%'],
]

ACKNOWLEDGMENT = ('Lời đầu tiên, nhóm nghiên cứu chúng em xin được gửi lời cảm ơn chân thành và sâu sắc nhất đến ThS. Trần Minh Toàn, '
    'giảng viên bộ môn Pháp Luật Kinh Doanh Bất Động Sản thuộc Khoa Luật, Trường Đại học Văn Lang. Trong suốt quá trình học tập '
    'và tiếp cận môn học, thầy đã không chỉ truyền đạt những kiến thức lý luận nền tảng quý báu về chế độ pháp lý kinh doanh bất '
    'động sản, điều kiện thành lập và hoạt động của các sàn giao dịch, mà còn luôn định hướng, khơi gợi tư duy pháp lý thực tiễn '
    'thông qua các tình huống giao dịch và tranh chấp điển hình trên thị trường. Chính những sự hướng dẫn tận tình và những bài '
    'giảng sinh động của thầy đã giúp nhóm có đủ cơ sở lý luận để hoàn thành bài báo cáo nghiên cứu về đề tài "Kinh doanh sàn '
    'giao dịch bất động sản: Những vấn đề lý luận và thực tiễn". Do điều kiện thời gian và năng lực nghiên cứu của sinh viên còn '
    'những giới hạn nhất định, bài báo cáo chắc chắn không tránh khỏi những thiếu sót. Nhóm chúng em rất mong nhận được những ý '
    'kiến đóng góp, nhận xét từ thầy để công trình nghiên cứu này được hoàn thiện và toàn diện hơn.')
DECLARATION = ('Nhóm nghiên cứu xin cam đoan bài báo cáo với đề tài "Kinh doanh sàn giao dịch bất động sản: Những vấn đề lý luận '
    'và thực tiễn" là công trình nghiên cứu độc lập của toàn thể các thành viên trong nhóm dưới sự định hướng của giảng viên phụ '
    'trách. Các nội dung phân tích, số liệu, quy định pháp luật và các tình huống thực tiễn được trình bày trong bài viết đều có '
    'nguồn gốc rõ ràng, được trích dẫn trung thực từ các văn bản quy phạm pháp luật hiện hành (như Luật Kinh doanh bất động sản, '
    'Luật Nhà ở và các văn bản hướng dẫn thi hành), các tài liệu khoa học pháp lý uy tín cũng như thực trạng hoạt động của các sàn '
    'giao dịch trên thị trường. Nhóm hoàn toàn chịu trách nhiệm trước Khoa Luật, Trường Đại học Văn Lang và pháp luật về tính trung '
    'thực, tính chính xác cũng như tính nguyên bản của nội dung bài báo cáo này.')

# Each block is one substantive page. Pages 3-12 are the ten report pages;
# the cover and the combined member-list/contents page are preliminary pages.
PAGES = [
    {
        'title': 'PHẦN MỞ ĐẦU',
        'sections': [
            ('1. Lý do chọn đề tài', [
                'Bất động sản có giá trị lớn, hồ sơ pháp lý và tình trạng thực tế gồm nhiều loại tài liệu. Người mua, người thuê không phải lúc nào cũng có điều kiện tự kiểm tra quyền định đoạt, quy hoạch, thế chấp, tranh chấp và điều kiện đưa tài sản vào kinh doanh. Sàn có thể tập trung thông tin, kết nối các bên, hỗ trợ đàm phán và lưu dấu giao dịch; vì vậy chất lượng hoạt động của sàn có ý nghĩa đối với tính minh bạch và phòng ngừa rủi ro.',
                'Tên gọi “sàn” không tự xác nhận mọi thông tin. Rủi ro xuất hiện khi hồ sơ không được đối chiếu, quảng cáo vượt tài liệu chứng minh, phí và phạm vi dịch vụ không rõ hoặc không lưu được người sửa tin. Giao dịch qua website cũng đặt ra yêu cầu xác thực danh tính, giữ lịch sử nội dung và chứng cứ điện tử. Luật Kinh doanh bất động sản năm 2023 cùng văn bản hướng dẫn quy định lại điều kiện, hoạt động và trách nhiệm của sàn; đề tài làm rõ các yêu cầu đó, các vướng mắc khi áp dụng và giải pháp khả thi.',
            ]),
            ('2. Mục tiêu và nhiệm vụ nghiên cứu', [
                'Báo cáo hướng đến việc hệ thống hóa cơ sở lý luận và pháp luật hiện hành, phân tích cách nghĩa vụ của sàn được thực hiện trong thực tiễn và kiến nghị biện pháp nâng cao minh bạch, xác định trách nhiệm, bảo vệ khách hàng. Nhiệm vụ cụ thể là làm rõ khái niệm, vai trò và sự khác nhau giữa sàn với môi giới; phân tích điều kiện, nội dung hoạt động, quyền và nghĩa vụ; đánh giá kiểm tra hồ sơ, thông tin môi giới, phí dịch vụ, giao dịch điện tử và giám sát; đề xuất giải pháp hoàn thiện và tổ chức thi hành. Không kết luận về mức độ phổ biến của vi phạm khi thiếu dữ liệu kiểm chứng; ví dụ không có nguồn được ghi là tình huống minh họa.',
            ]),
            ('3. Đối tượng, phạm vi và phương pháp', [
                'Đối tượng là quy định về dịch vụ sàn và quan hệ giữa sàn, doanh nghiệp kinh doanh bất động sản, môi giới, khách hàng và cơ quan quản lý. Phạm vi pháp lý lấy Luật Kinh doanh bất động sản năm 2023 cùng văn bản sửa đổi, hướng dẫn có hiệu lực tại ngày 12/10/2026, đối chiếu với pháp luật dân sự và giao dịch điện tử. Báo cáo không đi sâu vào đầu tư dự án hay tranh chấp quyền sử dụng đất thuần túy.',
                'Phương pháp gồm phân tích quy phạm, tổng hợp, so sánh khung luật mới với luật trước và nghiên cứu bản án có dẫn nguồn. Phần thực tiễn không được coi là khảo sát toàn quốc; ví dụ không có bản án ghi rõ là minh họa.',
            ]),
        ],
    },
    {
        'title': 'CHƯƠNG I. CƠ SỞ LÝ LUẬN VÀ PHÁP LUẬT HIỆN HÀNH',
        'sections': [
            ('1.1. Khái niệm, bản chất và vai trò', [
                'Khoản 10 Điều 3 Luật Kinh doanh bất động sản năm 2023 xác định sàn giao dịch bất động sản là nơi diễn ra các giao dịch mua bán, chuyển nhượng, cho thuê, cho thuê lại, cho thuê mua bất động sản được thành lập và hoạt động theo Luật. Sàn là tổ chức cung cấp một chuỗi dịch vụ: tiếp nhận và kiểm tra hồ sơ, niêm yết, cung cấp thông tin, hỗ trợ tìm kiếm và đàm phán, hỗ trợ lập hợp đồng, bàn giao, lưu trữ và xác nhận giao dịch theo phạm vi luật định. [1, khoản 10 Điều 3; Điều 56]',
                'Cần phân biệt sàn giao dịch với cá nhân môi giới. Môi giới là hoạt động trung gian giúp các bên tìm kiếm, đàm phán hoặc ký kết; sàn là doanh nghiệp hoặc hợp tác xã được thành lập, đăng ký và chịu trách nhiệm tổ chức hoạt động. Một website đăng tin không mặc nhiên là sàn theo Luật; phải xem xét mô hình, chức năng thực tế và việc đáp ứng điều kiện pháp luật. Người môi giới hành nghề phải đáp ứng điều kiện và làm việc trong doanh nghiệp kinh doanh dịch vụ sàn hoặc môi giới theo Điều 61. [1, khoản 11 Điều 3; Điều 61]',
                'Về bản chất, đây là quan hệ cung cấp dịch vụ được xác lập bằng thỏa thuận, đồng thời chịu sự quản lý chuyên ngành vì đối tượng giao dịch có giá trị lớn và thông tin không cân xứng. Sàn có thể giảm chi phí tìm kiếm, tạo một đầu mối lưu vết và giúp khách hàng đối chiếu tài liệu. Tuy vậy, việc sàn tiếp nhận hay niêm yết không thay thế giấy chứng nhận, xác nhận quy hoạch, quyết định của cơ quan có thẩm quyền hoặc việc thẩm tra độc lập của các bên; sàn cũng không tự động bảo đảm giá trị đầu tư hay khả năng sinh lời.',
                'Vai trò của sàn do đó gắn với ba chức năng: kiểm tra điều kiện trước khi đưa bất động sản lên sàn; truyền đạt thông tin đầy đủ, trung thực và có thể truy nguyên; hỗ trợ quy trình giao dịch và bảo quản hồ sơ. Nếu một khâu không được ghi nhận, việc xác định lỗi và quan hệ nhân quả khi tranh chấp xảy ra sẽ khó khăn hơn.',
            ]),
        ],
    },
    {
        'title': 'CHƯƠNG I. CƠ SỞ LÝ LUẬN VÀ PHÁP LUẬT HIỆN HÀNH',
        'sections': [
            ('1.2. Cơ sở pháp lý và thời điểm áp dụng', [
                'Khung pháp lý chính gồm Luật Kinh doanh bất động sản số 29/2023/QH15, Luật số 43/2024/QH15 sửa đổi thời điểm có hiệu lực, Nghị định số 96/2024/NĐ-CP và các sửa đổi liên quan tại Nghị định số 54/2026/NĐ-CP. Luật số 29 ban đầu ghi ngày hiệu lực 01/01/2025; khoản 1 Điều 82 được Luật số 43 sửa lại thành 01/08/2024. Vì vậy, khi xác định luật áp dụng phải dùng mốc đã được sửa đổi, không ghi mốc ban đầu như mốc hiện hành. [1; 2; 3; 4]',
                'Đối với dữ liệu, cần đọc cùng Nghị định số 357/2025/NĐ-CP về xây dựng và quản lý hệ thống thông tin, cơ sở dữ liệu về nhà ở và thị trường bất động sản. Đối với vi phạm hành chính, Nghị định số 339/2026/NĐ-CP quy định xử phạt trong các lĩnh vực có kinh doanh bất động sản và có hiệu lực từ 26/08/2026. Hai nghị định này hỗ trợ phân tích quản trị dữ liệu và cơ chế xử lý; báo cáo không nêu mức phạt cụ thể để tránh dùng sai trường hợp hoặc điều khoản. [5; 6]',
            ]),
            ('1.3. Điều kiện thành lập, đăng ký và hoạt động', [
                'Tổ chức, cá nhân kinh doanh dịch vụ sàn phải thành lập doanh nghiệp hoặc hợp tác xã, liên hiệp hợp tác xã theo khoản 5 Điều 9. Tên doanh nghiệp phải có cụm từ “sàn giao dịch bất động sản”. Trước khi hoạt động, sàn gửi hồ sơ đến cơ quan quản lý nhà nước về kinh doanh bất động sản cấp tỉnh nơi đặt trụ sở để được cấp giấy phép hoạt động. Do đó, chỉ đăng ký doanh nghiệp hoặc tạo website chưa đủ để chứng minh một tổ chức đã đủ điều kiện hoạt động như sàn theo Luật. [1, khoản 5 Điều 9; Điều 54]',
                'Người quản lý, điều hành phải có quyền quản lý doanh nghiệp tại Việt Nam và hoàn thành khóa đào tạo quản lý, điều hành sàn. Sàn phải đăng ký địa điểm hoạt động cố định, có cơ sở vật chất và kỹ thuật phù hợp, ban hành và công khai quy trình giao dịch. Nhân viên môi giới làm việc trong sàn phải đáp ứng điều kiện hành nghề; doanh nghiệp còn có nghĩa vụ quản lý, bồi dưỡng nhân viên. Các yêu cầu này tạo cơ sở xác định trách nhiệm nội bộ, nhưng không làm mất nghĩa vụ kiểm tra riêng của sàn với từng hồ sơ. [1, Điều 55, khoản 7–9 Điều 58]',
                'Về đăng ký và vận hành, doanh nghiệp cần lưu hồ sơ thành lập, giấy phép, quy chế, quy trình, hồ sơ bất động sản, tài liệu niêm yết và chứng cứ cập nhật. Khi kiểm tra một tin đăng, cần xác định đúng loại tài sản và loại giao dịch để áp dụng đúng điều kiện pháp luật; giấy tờ cần có cho nhà ở hình thành trong tương lai không đồng nhất với nhà ở có sẵn hay quyền sử dụng đất trong dự án. [1, Điều 6; Điều 24 và Điều 55; 3; 4]',
            ]),
        ],
    },
    {
        'title': 'CHƯƠNG I. CƠ SỞ LÝ LUẬN VÀ PHÁP LUẬT HIỆN HÀNH',
        'sections': [
            ('1.4. Nội dung hoạt động, quyền và nghĩa vụ của sàn', [
                'Điều 56 yêu cầu sàn kiểm tra tính pháp lý và điều kiện đưa bất động sản vào kinh doanh trước khi giao dịch qua sàn; niêm yết và cung cấp thông tin về bất động sản đủ điều kiện. Sàn còn hỗ trợ khách hàng tìm kiếm, lựa chọn; thực hiện môi giới và hỗ trợ đàm phán; hỗ trợ lập, ký hợp đồng, thanh toán, bàn giao; lưu giữ hồ sơ; xác nhận giao dịch và cung cấp thông tin giao dịch cho cơ quan quản lý cấp tỉnh. [1, Điều 56]',
                'Quyền của sàn gồm yêu cầu khách hàng cung cấp hồ sơ, từ chối tài sản không đủ điều kiện, thu phí theo thỏa thuận và yêu cầu bồi thường nếu khách hàng có lỗi gây thiệt hại. Nghĩa vụ tương ứng là bảo đảm điều kiện trước khi đưa tài sản lên sàn; niêm yết, cung cấp hồ sơ đầy đủ và trung thực; duy trì điều kiện hoạt động; báo cáo, chịu kiểm tra; thực hiện nghĩa vụ thuế; bồi thường thiệt hại do lỗi của sàn; quản lý và bồi dưỡng nhân viên; lưu trữ hồ sơ. [1, Điều 57–58]',
                'Quy trình thực hiện nên bắt đầu bằng tiếp nhận và lập danh mục tài liệu; đối chiếu chủ thể cung cấp với giấy tờ về quyền và quyền đưa tài sản ra giao dịch; kiểm tra tình trạng thế chấp, tranh chấp, điều kiện dự án theo loại bất động sản; ghi ngày kiểm tra và nguồn tra cứu; sau đó mới phê duyệt tin đăng. Nếu hồ sơ thiếu, có mâu thuẫn hoặc tình trạng pháp lý thay đổi, sàn cần tạm dừng hoặc cập nhật niêm yết, lưu lại phiên bản và thông báo cho khách hàng. Đây là cách biến nghĩa vụ pháp lý thành dấu vết có thể kiểm tra.',
                'Về thanh toán, Điều 48 phân biệt hai quan hệ. Các bên giao dịch bất động sản thỏa thuận phương thức thanh toán trong hợp đồng và tuân thủ pháp luật. Riêng doanh nghiệp kinh doanh bất động sản hoặc dịch vụ bất động sản nhận tiền thanh toán từ khách hàng theo hợp đồng kinh doanh hoặc hợp đồng dịch vụ phải nhận qua tài khoản tại tổ chức tín dụng hợp pháp. Chính sách khuyến khích thanh toán không dùng tiền mặt không có nghĩa mọi giao dịch bất động sản đều bắt buộc phải thực hiện qua sàn. [1, khoản 1–2 Điều 48; khoản 6–7 Điều 7]',
            ]),
        ],
    },
    {
        'title': 'CHƯƠNG II. THỰC TIỄN ÁP DỤNG VÀ BẤT CẬP',
        'sections': [
            ('2.1. Kiểm tra hồ sơ và minh bạch thông tin', [
                'Thông tin ban đầu thường do chủ sở hữu, chủ đầu tư hoặc người được ủy quyền cung cấp. Sàn không thể chỉ sao chép nội dung quảng cáo rồi xem đó là kết quả kiểm tra. Tùy loại tài sản, cần đối chiếu giấy tờ quyền sở hữu hoặc quyền sử dụng, căn cứ đại diện, tình trạng thế chấp và tranh chấp, quyết định hoặc hồ sơ dự án, tình trạng xây dựng, điều kiện bán hoặc cho thuê mua khi pháp luật yêu cầu. Các thông tin công khai theo Điều 6 phải đầy đủ, trung thực, chính xác và được cập nhật khi thay đổi. [1, Điều 6; Điều 56; Điều 58]',
                'Tình huống minh họa: bên bán đề nghị đăng một căn nhà với thông tin “đủ giấy tờ”, nhưng hồ sơ thể hiện tài sản đang thế chấp và người ký chưa chứng minh được quyền đại diện. Sàn cần dừng niêm yết phần khẳng định chưa được chứng minh, yêu cầu tài liệu bổ sung và ghi rõ trạng thái xác minh. Nếu vẫn đăng nội dung chắc chắn rằng tài sản không bị hạn chế, khách hàng có thể dựa vào thông tin đó để đặt cọc và phát sinh tranh chấp. Đây là ví dụ phân tích, không phải vụ việc có thật.',
                'Khâu khó nằm ở độ sâu và nguồn của việc kiểm tra. Một bản chụp giấy tờ do bên bán gửi không tự chứng minh giấy tờ còn hiệu lực hoặc phản ánh đầy đủ tình trạng tại ngày giao dịch. Sàn nên ghi lại bản tài liệu đã nhận, người cung cấp, ngày tra cứu, kết quả đối chiếu và giới hạn kiểm tra; thông tin chưa độc lập xác nhận cần được thể hiện đúng trạng thái, không biến thành lời bảo đảm. Khách hàng cũng cần được hướng dẫn kiểm tra tại cơ quan có thẩm quyền đối với vấn đề mà sàn không thể xác nhận. Đối với dự án, việc kiểm tra có thể bao gồm quyết định đầu tư, hồ sơ đất, quy hoạch được duyệt, tiến độ xây dựng, tình trạng thế chấp và văn bản xác nhận đủ điều kiện bán nếu luật yêu cầu. Danh mục cụ thể phải tùy loại bất động sản; không nên dùng một bộ giấy tờ chung cho mọi giao dịch. [1, Điều 6; Điều 24]',
            ]),
        ],
    },
    {
        'title': 'CHƯƠNG II. THỰC TIỄN ÁP DỤNG VÀ BẤT CẬP',
        'sections': [
            ('2.2. Môi giới, thông tin sai lệch và tình huống xét xử', [
                'Nguy cơ khác xuất hiện khi nhân viên mô tả tài sản là “sắp tăng giá”, “đã có nhiều người đặt chỗ” hoặc “đủ điều kiện chuyển nhượng” nhưng không chỉ ra căn cứ. Một nhận định thị trường hay nội dung quảng cáo không đương nhiên là vi phạm; cần xem chính xác thông tin nào được đưa ra, có chủ ý sai lệch hay không, khách hàng đã dựa vào thông tin đó thế nào, thiệt hại nào phát sinh và ai chịu trách nhiệm. Luật cấm giả mạo, cố ý làm sai lệch thông tin và gian lận; sàn có nghĩa vụ quản lý nhân viên và chịu trách nhiệm với thông tin do mình niêm yết, cung cấp. [1, khoản 2, 4 Điều 8; khoản 2, 6–8 Điều 58]',
                'Nếu bất động sản chưa đủ điều kiện giao dịch, khách hàng có thể gặp rủi ro không ký được hợp đồng, chậm bàn giao hoặc tranh chấp tiền đặt cọc. Sàn phải kiểm tra điều kiện trước khi đưa tài sản lên sàn; khi hồ sơ thiếu hoặc mâu thuẫn cần dừng hay giới hạn nội dung niêm yết và yêu cầu làm rõ. Trách nhiệm của chủ sở hữu, chủ đầu tư, môi giới và sàn phải xác định theo nghĩa vụ, hành vi, lỗi, thiệt hại và quan hệ nhân quả; không thể mặc nhiên quy toàn bộ thiệt hại cho một chủ thể chỉ vì họ tham gia giao dịch. [1, Điều 56; 7]',
                'Tình huống có thật được phản ánh trong Bản án số 325/2024/HS-PT ngày 26/07/2024 của Tòa án nhân dân cấp cao tại Đà Nẵng: bị cáo đưa thông tin sai về việc doanh nghiệp có quyền tham gia môi giới dự án, nhận tiền đặt chỗ hoặc góp vốn của khách hàng, không thực hiện như đã hứa và bị kết tội lừa đảo chiếm đoạt tài sản. Vụ án cho thấy cần xác minh tư cách, hợp đồng ủy quyền hoặc phân phối, người nhận tiền và chứng từ. Đây là vụ án hình sự về hành vi cụ thể, xảy ra trước khi Luật số 29/2023/QH15 có hiệu lực; không phải căn cứ quy trách nhiệm chung cho mọi sàn hoặc cách giải thích các quy định mới. [9]',
                'Tranh chấp phí dịch vụ thường xoay quanh mức phí, bên phải trả, thời điểm phát sinh, điều kiện được hưởng phí và xử lý khi giao dịch không thành. Hợp đồng nên nêu công việc sàn thực hiện, tài liệu bàn giao, cách tính phí, thời điểm thanh toán, khoản nào được hoàn lại và phương thức giải quyết khiếu nại. Nếu chỉ thỏa thuận miệng, các bên có thể bất đồng về việc môi giới đã hoàn thành hay chưa và ai là người trả phí; vì vậy cần lưu hợp đồng, tin nhắn, lịch xem tài sản và biên bản giao nhận. Việc nhận tiền dịch vụ phải tuân thủ khoản 2 Điều 48; tiền giao dịch giữa bên mua và bên bán cần ghi nhận riêng, tránh nhập nhằng với phí môi giới. [1, khoản 2 Điều 48; Điều 63]',
            ]),
        ],
    },
    {
        'title': 'CHƯƠNG II. THỰC TIỄN ÁP DỤNG VÀ BẤT CẬP',
        'sections': [
            ('2.3. Giao dịch trực tuyến và chứng cứ điện tử', [
                'Sàn có thể tổ chức giao dịch trực tiếp hoặc điện tử. Khi dùng phương thức điện tử, quy trình cần xác thực tài khoản, người đại diện, thời điểm đăng hoặc sửa thông tin, phiên bản hợp đồng, việc khách hàng chấp thuận và người nhận khoản tiền. Nhật ký truy cập, email, tin nhắn, tệp đính kèm và biên nhận phải được lưu theo cách có thể truy nguyên. Luật Kinh doanh bất động sản yêu cầu giao dịch điện tử tuân thủ Luật Giao dịch điện tử và quy định liên quan. [1, khoản 5 Điều 53; 8]',
                'Tình huống minh họa: khách hàng nhận được ảnh chụp tin đăng có giá và diện tích khác với nội dung trên website sau khi có tranh chấp. Nếu sàn không lưu lịch sử chỉnh sửa hoặc không xác định tài khoản đã cập nhật tin, rất khó chứng minh thông tin nào được xem tại thời điểm khách hàng đặt cọc. Biện pháp thực tế là lưu dấu thời gian, tài khoản thao tác, phiên bản nội dung, trạng thái xác minh và thông báo thay đổi; đồng thời bảo vệ dữ liệu cá nhân, giới hạn quyền truy cập và thời hạn lưu theo quy định.',
                'Cần phân biệt nền tảng đăng tin với sàn giao dịch. Một nền tảng có thể chỉ cung cấp không gian đăng tin; một sàn theo Luật có các chức năng, điều kiện cấp phép và nghĩa vụ riêng. Đánh giá phải dựa trên mô hình thực tế và hành vi cung cấp dịch vụ, không chỉ dựa vào tên gọi hoặc việc có ứng dụng. Đối với giao dịch trực tuyến, việc có tin nhắn hay chuyển khoản cũng không tự chứng minh tài sản đủ điều kiện hoặc thỏa thuận đã đáp ứng mọi điều kiện hình thức của hợp đồng. Việc thu thập và lưu trữ dữ liệu khách hàng còn đòi hỏi xác định mục đích, người được truy cập, thời hạn lưu và cách xử lý khi có yêu cầu chỉnh sửa hoặc xóa. Khi phát sinh tranh chấp, bản chụp màn hình riêng lẻ có thể không cho biết ai đăng, nội dung đã bị sửa lúc nào hay phiên bản nào được các bên chấp thuận; vì vậy nhật ký hệ thống và bản lưu có thời điểm là chứng cứ quan trọng. [8]',
            ]),
        ],
    },
    {
        'title': 'CHƯƠNG II. THỰC TIỄN ÁP DỤNG VÀ BẤT CẬP',
        'sections': [
            ('2.4. Giám sát, phối hợp và các điểm nghẽn thực thi', [
                'Cơ quan quản lý cấp tỉnh được giao quản lý, giám sát, kiểm tra định kỳ hoặc đột xuất, cấp hoặc thu hồi giấy phép và quyết định tạm ngừng hay chấm dứt hoạt động theo luật. Một giao dịch có thể liên quan đến chủ đầu tư, chủ sở hữu, doanh nghiệp môi giới, sàn và cá nhân môi giới; tài liệu lại nằm ở nhiều hệ thống. Việc xác định người chịu trách nhiệm cần dựa trên vai trò, hồ sơ, tin đăng, hợp đồng, dòng tiền và hành vi cụ thể, không chỉ dựa vào việc họ cùng xuất hiện trong giao dịch. [1, Điều 60]',
                'Nghị định số 357/2025/NĐ-CP tạo khung quản lý hệ thống thông tin và cơ sở dữ liệu về nhà ở, thị trường bất động sản; Nghị định số 339/2026/NĐ-CP quy định xử phạt vi phạm hành chính trong lĩnh vực có liên quan. Hiệu quả phụ thuộc vào chất lượng và khả năng cập nhật, đối chiếu dữ liệu, cũng như chứng cứ đủ để xác định hành vi. Tin đăng trực tuyến có thể bị sửa, xóa hoặc sao chép trên nhiều nền tảng, làm khó việc xác định phiên bản và người thao tác. Các quy định cũng giao cho doanh nghiệp thêm công việc về hồ sơ, báo cáo, lưu trữ và công nghệ; mức độ khó khăn cần được đánh giá bằng khảo sát hoặc báo cáo, không suy đoán cho mọi sàn. [5; 6]',
                'Các điểm cần kiểm soát gồm hồ sơ đầu vào thiếu hoặc lỗi thời; quảng cáo vượt tài liệu chứng minh; quy trình không giữ dấu vết; và hợp đồng không rõ phí, công việc, trách nhiệm. Việc áp dụng đồng thời luật về kinh doanh bất động sản, đất đai, nhà ở, xây dựng, dân sự và giao dịch điện tử có thể đòi hỏi đối chiếu nhiều nguồn. Đây là nguy cơ và vấn đề quản trị, không tự chứng minh rằng hệ thống pháp luật mâu thuẫn hoặc mọi doanh nghiệp vi phạm; kết luận về mức độ phổ biến phải dựa trên số liệu, quyết định xử lý hoặc tập hợp bản án.',
                'Việc áp dụng quy định mới còn đòi hỏi sàn rà soát điều kiện hoạt động, quy chế, quy trình giao dịch, nhân sự và hệ thống lưu trữ. Doanh nghiệp có thể cần đầu tư thêm thời gian, nhân lực và công nghệ để xử lý hồ sơ điện tử, cập nhật dữ liệu và đào tạo nhân viên; mức độ đáp ứng khác nhau giữa các đơn vị. Muốn kết luận đây là khó khăn phổ biến cần có báo cáo quản lý, kết quả khảo sát doanh nghiệp hoặc dữ liệu triển khai cụ thể, thay vì khái quát từ một vài tình huống.',
            ]),
        ],
    },
    {
        'title': 'CHƯƠNG III. KIẾN NGHỊ HOÀN THIỆN',
        'sections': [
            ('3.1. Hoàn thiện hướng dẫn và quy trình kiểm tra', [
                'Cơ quan có thẩm quyền nên hướng dẫn danh mục hồ sơ theo loại bất động sản và giao dịch, kèm nguồn tra cứu, thời điểm đối chiếu và cách ghi nhận phần chưa xác minh. Hướng dẫn cần phân biệt việc sàn kiểm tra với việc cơ quan nhà nước xác nhận và nghĩa vụ khách hàng cung cấp tài liệu. Biểu mẫu nên ghi tài liệu đã nhận, người cung cấp, kết quả, phần còn thiếu, ngày cập nhật và quyết định niêm yết hay từ chối.',
                'Cần làm rõ cách xác định trách nhiệm khi thông tin sai do khách hàng hoặc môi giới cung cấp, khi sàn biết hoặc phải biết sai lệch, hoặc khi nội dung sửa đổi không được thông báo. Bồi thường phải căn cứ pháp luật, lỗi, thiệt hại và quan hệ nhân quả; sàn không mặc nhiên bảo lãnh mọi giao dịch. Cơ quan quản lý nên công khai thủ tục cấp phép, tạm ngừng, thu hồi và kênh tra cứu trạng thái sàn.',
            ]),
            ('3.2. Nâng cao hiệu quả tổ chức và quản lý dữ liệu', [
                'Mỗi sàn nên phân công trách nhiệm từ tiếp nhận hồ sơ, kiểm tra điều kiện, duyệt tin và cập nhật đến lưu hợp đồng, giao dịch, khiếu nại. Nhân viên môi giới không nên cam kết về lợi nhuận hoặc tính pháp lý nếu thiếu chứng cứ; quản lý cần duyệt và lưu căn cứ cho nội dung quảng cáo có yếu tố pháp lý.',
                'Hệ thống dữ liệu nên lưu nguồn chính thức, lịch sử sửa đổi, tài khoản thao tác và phân quyền truy cập. Việc kết nối, chia sẻ phải đúng thẩm quyền, không tự đặt nghĩa vụ cập nhật “thời gian thực” nếu luật không quy định. Cần bảo vệ dữ liệu cá nhân và có phương án khôi phục chứng cứ khi hệ thống gặp sự cố. [5; 8]',
            ]),
            ('3.3. Giám sát, chế tài và bảo vệ khách hàng', [
                'Cơ quan cấp tỉnh nên kết hợp kiểm tra định kỳ với lựa chọn theo rủi ro: chưa có giấy phép, hồ sơ thiếu căn cứ, tin quảng cáo mâu thuẫn hoặc phản ánh lặp lại. Phối hợp liên ngành cần rõ đầu mối, loại dữ liệu được chia sẻ và quy trình chuyển hồ sơ. Khi xử lý, phải đối chiếu đúng hành vi với điều khoản còn hiệu lực, áp dụng biện pháp khắc phục và theo dõi kết quả.',
                'Khách hàng nên xem hợp đồng dịch vụ, hồ sơ tài sản, phạm vi sàn đã kiểm tra, bên nhận phí và điều kiện hoàn phí; lưu tin đăng, thư từ, hợp đồng, chứng từ và biên bản bàn giao. Sàn cần công khai đầu mối khiếu nại và thời hạn phản hồi. Các bước này giúp trách nhiệm giữa các bên rõ và có chứng cứ hơn.',
                'Sàn nên hướng dẫn nghiệp vụ định kỳ cho người quản lý và môi giới về kiểm tra hồ sơ, giới hạn của quảng cáo, bảo vệ dữ liệu và tiếp nhận khiếu nại. Hiệp hội nghề nghiệp có thể hỗ trợ tài liệu thực hành và chuẩn mực ứng xử; đây là giải pháp khuyến nghị nâng chất lượng dịch vụ, không thay thế điều kiện hay nghĩa vụ pháp luật quy định.',
            ]),
        ],
    },
    {
        'title': 'KẾT LUẬN',
        'sections': [
            (None, [
                'Sàn giao dịch bất động sản giữ vai trò tổ chức dịch vụ, hỗ trợ thông tin và kết nối các bên; vai trò đó đi kèm nghĩa vụ kiểm tra điều kiện, cung cấp thông tin đầy đủ và trung thực, lưu trữ hồ sơ, quản lý nhân viên và bồi thường thiệt hại do lỗi của mình. Luật Kinh doanh bất động sản số 29/2023/QH15, được sửa đổi về thời điểm hiệu lực bởi Luật số 43/2024/QH15, cùng các nghị định hướng dẫn và văn bản về dữ liệu, xử phạt tạo thành khung pháp lý cần được áp dụng đồng bộ. [1–6]',
                'Phân tích lý luận và tình huống thực tiễn cho thấy rủi ro tập trung ở việc kiểm tra tài liệu đầu vào, xác định người có quyền đưa tài sản ra giao dịch, quản lý lời giới thiệu của môi giới, xác định phí và lưu chứng cứ điện tử. Bản án được dẫn chỉ minh họa hậu quả của một hành vi cụ thể, không chứng minh mọi sàn đều có cùng vấn đề. Muốn đánh giá thực trạng phổ biến cần có số liệu, báo cáo quản lý hoặc tập hợp bản án có phương pháp.',
                'Để nâng cao hiệu quả, cần cụ thể hóa hồ sơ kiểm tra, chuẩn hóa hợp đồng và quy trình nội bộ, giữ dấu vết cập nhật, kết nối dữ liệu theo đúng thẩm quyền, tăng cường giám sát dựa trên rủi ro và nâng cao hiểu biết pháp lý của khách hàng. Khi mỗi chủ thể thực hiện đúng nghĩa vụ và có thể chứng minh việc đã làm, sàn mới góp phần thiết thực vào sự minh bạch, an toàn và phát triển bền vững của thị trường bất động sản.',
            ]),
            ('DANH MỤC TÀI LIỆU THAM KHẢO', [
                '[1] Quốc hội, Luật Kinh doanh bất động sản số 29/2023/QH15 ngày 28/11/2023, đặc biệt các Điều 3, 6–9, 48, 53–61.',
                '[2] Quốc hội, Luật số 43/2024/QH15 ngày 29/06/2024, Điều 3 về thời điểm có hiệu lực của Luật Kinh doanh bất động sản.',
                '[3] Chính phủ, Nghị định số 96/2024/NĐ-CP ngày 24/07/2024 quy định chi tiết một số điều của Luật Kinh doanh bất động sản.',
                '[4] Chính phủ, Nghị định số 54/2026/NĐ-CP ngày 09/02/2026 sửa đổi, bổ sung một số nghị định trong lĩnh vực nhà ở, kinh doanh bất động sản.',
                '[5] Chính phủ, Nghị định số 357/2025/NĐ-CP ngày 31/12/2025 về hệ thống thông tin, cơ sở dữ liệu nhà ở và thị trường bất động sản.',
                '[6] Chính phủ, Nghị định số 339/2026/NĐ-CP ngày 26/08/2026 về xử phạt vi phạm hành chính trong lĩnh vực xây dựng, nhà ở và kinh doanh bất động sản.',
                '[7] Quốc hội, Bộ luật Dân sự số 91/2015/QH13 ngày 24/11/2015.',
                '[8] Quốc hội, Luật Giao dịch điện tử số 20/2023/QH15 ngày 22/06/2023.',
                '[9] Tòa án nhân dân cấp cao tại Đà Nẵng, Bản án số 325/2024/HS-PT ngày 26/07/2024.',
            ]),
        ],
    },
]

TOC = [
    (0, 'LỜI CẢM ƠN', 2),
    (0, 'LỜI CAM ĐOAN', 2),
    (0, 'PHẦN MỞ ĐẦU', 3),
    (0, 'CHƯƠNG I. CƠ SỞ LÝ LUẬN VÀ PHÁP LUẬT HIỆN HÀNH', 4),
    (1, '1.1. Khái niệm, bản chất và vai trò', 4),
    (1, '1.2. Cơ sở pháp lý và thời điểm áp dụng', 5),
    (1, '1.3. Điều kiện thành lập, đăng ký và hoạt động', 5),
    (1, '1.4. Nội dung, quyền và nghĩa vụ của sàn', 6),
    (0, 'CHƯƠNG II. THỰC TIỄN ÁP DỤNG VÀ BẤT CẬP', 7),
    (1, '2.1. Kiểm tra hồ sơ và minh bạch thông tin', 7),
    (1, '2.2. Môi giới, thông tin sai lệch và tình huống xét xử', 8),
    (1, '2.3. Giao dịch trực tuyến và chứng cứ điện tử', 9),
    (1, '2.4. Giám sát, phối hợp và điểm nghẽn thực thi', 10),
    (0, 'CHƯƠNG III. KIẾN NGHỊ HOÀN THIỆN', 11),
    (0, 'KẾT LUẬN VÀ DANH MỤC TÀI LIỆU THAM KHẢO', 12),
]


def set_run_font(run, size=13, bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn('w:' + key), 'Times New Roman')


def add_page_field(paragraph):
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0 if level == 1 else 3)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 2)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), 14 if level == 1 else 13, True)
    return p


def add_body(doc, text, reference=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if reference else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(1 if reference else 4)
    p.paragraph_format.line_spacing = 1.3
    if reference:
        p.paragraph_format.left_indent = Cm(.55)
        p.paragraph_format.first_line_indent = Cm(-.55)
    else:
        p.paragraph_format.first_line_indent = Cm(.65)
    set_run_font(p.add_run(text), 13)
    return p


def add_frontmatter(doc, title, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(title), 11, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Cm(.4)
    set_run_font(p.add_run(text), 10)


def set_cell(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run(text), 10 if not bold else 10, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3.5), Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
    sec.footer_distance = Cm(.8)
    sec.different_first_page_header_footer = True
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'; normal.font.size = Pt(13)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.paragraph_format.line_spacing = 1.3
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)
    sec.first_page_footer.paragraphs[0].text = ''

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    if LOGO.exists(): p.add_run().add_picture(str(LOGO), width=Inches(.55))
    for text in ('TRƯỜNG ĐẠI HỌC VĂN LANG', 'KHOA LUẬT'):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(text), 13, True)
    for _ in range(3):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    for text, size, after in [
        ('BÀI TẬP NHÓM SỐ 1', 15, 12),
        ('MÔN PHÁP LUẬT KINH DOANH BẤT ĐỘNG SẢN', 14, 12),
        ('KINH DOANH DỊCH VỤ SÀN GIAO DỊCH BẤT ĐỘNG SẢN', 14, 3),
        ('NHỮNG VẤN ĐỀ LÝ LUẬN VÀ THỰC TIỄN', 14, 20),
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        set_run_font(p.add_run(text), size, True)
    for text in (
        'Nhóm thực hiện: Nhóm 5',
        'Mã lớp học phần: 261_71LAWS40482_01',
        'Giảng viên hướng dẫn: ThS. Trần Minh Toàn',
    ):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_after = Pt(6)
        set_run_font(p.add_run(text), 13)
    for _ in range(4):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('Thành phố Hồ Chí Minh, ngày 12 tháng 10 năm 2026'), 13)
    doc.add_page_break()

    add_heading(doc, 'DANH SÁCH THÀNH VIÊN')
    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.autofit = False
    widths = [Inches(.4), Inches(1.38), Inches(1.25), Inches(1.85), Inches(.65)]
    headers = ['STT', 'Họ và tên', 'MSSV', 'Phân công', 'Đánh giá']
    for c, text, width in zip(tbl.rows[0].cells, headers, widths):
        c.width = width; set_cell(c, text, WD_ALIGN_PARAGRAPH.CENTER, True)
    for row in MEMBERS:
        cells = tbl.add_row().cells
        for i, (c, text, width) in enumerate(zip(cells, row, widths)):
            c.width = width
            set_cell(c, text, WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2, 4) else WD_ALIGN_PARAGRAPH.LEFT)
    add_frontmatter(doc, 'LỜI CẢM ƠN', ACKNOWLEDGMENT)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run('Nhóm chúng em xin trân trọng cảm ơn!'), 10)
    add_frontmatter(doc, 'LỜI CAM ĐOAN', DECLARATION)
    add_heading(doc, 'MỤC LỤC')
    for level, label, page in TOC:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Inches(.18 if level else 0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(5.95), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        set_run_font(p.add_run(label), 10 if level else 10, level == 0)
        set_run_font(p.add_run('\t' + str(page)), 10)
    doc.add_page_break()

    for page in PAGES:
        add_heading(doc, page['title'])
        for sub, paragraphs in page['sections']:
            if sub:
                add_heading(doc, sub, 2)
            for text in paragraphs:
                add_body(doc, text, reference=(page['title'] == 'KẾT LUẬN' and sub == 'DANH MỤC TÀI LIỆU THAM KHẢO'))
        doc.add_page_break()
    # Remove the final empty page caused by the trailing break.
    if doc.paragraphs and doc.paragraphs[-1].text == '':
        p = doc.paragraphs[-1]._element
        p.getparent().remove(p)
    doc.core_properties.title = 'Kinh doanh dịch vụ sàn giao dịch bất động sản - Nhóm 5'
    doc.core_properties.subject = 'Báo cáo Khoa Luật - 10 trang nội dung'
    doc.core_properties.author = 'Nhóm 5'
    doc.save(FULL_DOCX)


font_dir = Path(r'C:\Windows\Fonts')
for name, file in [
    ('TNR', 'times.ttf'), ('TNR-Bold', 'timesbd.ttf'),
    ('TNR-Italic', 'timesi.ttf'), ('TNR-BoldItalic', 'timesbi.ttf'),
]:
    pdfmetrics.registerFont(TTFont(name, str(font_dir / file)))
pdfmetrics.registerFontFamily('TNR', normal='TNR', bold='TNR-Bold', italic='TNR-Italic', boldItalic='TNR-BoldItalic')

ST = {
    'body': ParagraphStyle('Body', fontName='TNR', fontSize=13, leading=17,
                           alignment=TA_JUSTIFY, firstLineIndent=.65*cm, spaceAfter=4,
                           textColor=colors.black),
    'ref': ParagraphStyle('Ref', fontName='TNR', fontSize=13, leading=17,
                          alignment=TA_LEFT, leftIndent=.55*cm, firstLineIndent=-.55*cm,
                          spaceAfter=3, textColor=colors.black),
    'head': ParagraphStyle('Head', fontName='TNR-Bold', fontSize=14, leading=18,
                           alignment=TA_CENTER, spaceAfter=5, keepWithNext=True,
                           textColor=colors.black),
    'sub': ParagraphStyle('Sub', fontName='TNR-Bold', fontSize=13, leading=16,
                          alignment=TA_LEFT, spaceBefore=3, spaceAfter=2,
                          keepWithNext=True, textColor=colors.black),
    'cover': ParagraphStyle('Cover', fontName='TNR', fontSize=13, leading=18,
                            alignment=TA_CENTER, spaceAfter=5, textColor=colors.black),
    'coverbold': ParagraphStyle('CoverBold', fontName='TNR-Bold', fontSize=14, leading=18,
                                alignment=TA_CENTER, spaceAfter=4, textColor=colors.black),
    'smallbold': ParagraphStyle('SmallBold', fontName='TNR-Bold', fontSize=13, leading=16,
                                alignment=TA_CENTER, spaceAfter=4, textColor=colors.black),
    'table': ParagraphStyle('Table', fontName='TNR', fontSize=9, leading=11,
                            alignment=TA_LEFT, textColor=colors.black),
    'tablec': ParagraphStyle('TableC', fontName='TNR', fontSize=9, leading=11,
                             alignment=TA_CENTER, textColor=colors.black),
    'tableb': ParagraphStyle('TableB', fontName='TNR-Bold', fontSize=9, leading=11,
                             alignment=TA_CENTER, textColor=colors.black),
    'toc': ParagraphStyle('TOC', fontName='TNR', fontSize=10, leading=12,
                          spaceAfter=0, textColor=colors.black),
    'tocb': ParagraphStyle('TOCB', fontName='TNR-Bold', fontSize=10, leading=12,
                           spaceAfter=0, textColor=colors.black),
    'front': ParagraphStyle('Front', fontName='TNR', fontSize=10, leading=12,
                            alignment=TA_JUSTIFY, firstLineIndent=.4*cm, spaceAfter=1,
                            textColor=colors.black),
    'fronthead': ParagraphStyle('FrontHead', fontName='TNR-Bold', fontSize=11, leading=13,
                                alignment=TA_CENTER, spaceBefore=2, spaceAfter=1,
                                textColor=colors.black),
}


def para(text, style='body'):
    return Paragraph(escape(text), ST[style])


def pdf_head(text, level=1):
    return para(text, 'head' if level == 1 else 'sub')


def make_cover():
    logo = Image(str(LOGO), width=.5*inch, height=.6*inch) if LOGO.exists() else Spacer(1, .6*inch)
    logo.hAlign = 'CENTER'
    out = [logo, Spacer(1, .10*inch)]
    out += [para('TRƯỜNG ĐẠI HỌC VĂN LANG', 'coverbold'), para('KHOA LUẬT', 'coverbold')]
    out += [Spacer(1, .65*inch), para('BÀI TẬP NHÓM SỐ 1', 'coverbold'),
            para('MÔN PHÁP LUẬT KINH DOANH BẤT ĐỘNG SẢN', 'coverbold'), Spacer(1, .16*inch),
            para('KINH DOANH DỊCH VỤ SÀN GIAO DỊCH BẤT ĐỘNG SẢN', 'coverbold'),
            para('NHỮNG VẤN ĐỀ LÝ LUẬN VÀ THỰC TIỄN', 'coverbold'), Spacer(1, .16*inch)]
    for line in ('Nhóm thực hiện: Nhóm 5', 'Mã lớp học phần: 261_71LAWS40482_01',
                 'Giảng viên hướng dẫn: ThS. Trần Minh Toàn'):
        out.append(para(line, 'cover'))
    out += [Spacer(1, .30*inch), para('Thành phố Hồ Chí Minh, ngày 12 tháng 10 năm 2026', 'cover')]
    return out


def make_members_toc():
    out = [pdf_head('DANH SÁCH THÀNH VIÊN')]
    data = [[para(x, 'tableb') for x in ['STT', 'Họ và tên', 'MSSV', 'Phân công', 'Đánh giá']]]
    for row in MEMBERS:
        data.append([para(row[0], 'tablec'), para(row[1], 'table'), para(row[2], 'tablec'),
                     para(row[3], 'table'), para(row[4], 'tablec')])
    tab = Table(data, colWidths=[.4*inch, 1.4*inch, 1.28*inch, 1.78*inch, .65*inch], repeatRows=1, hAlign='CENTER')
    tab.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), .45, colors.black), ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('LEFTPADDING', (0,0), (-1,-1), 3), ('RIGHTPADDING', (0,0), (-1,-1), 3),
        ('TOPPADDING', (0,0), (-1,-1), 3), ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('BACKGROUND', (0,0), (-1,-1), colors.white),
    ]))
    out += [tab, Spacer(1, .05*inch), para('LỜI CẢM ƠN', 'fronthead'), para(ACKNOWLEDGMENT, 'front'),
            Paragraph(escape('Nhóm chúng em xin trân trọng cảm ơn!'), ParagraphStyle('Thanks', parent=ST['front'], alignment=TA_RIGHT)),
            para('LỜI CAM ĐOAN', 'fronthead'), para(DECLARATION, 'front'), Spacer(1, .04*inch), pdf_head('MỤC LỤC')]
    rows = []
    for level, label, page in TOC:
        style = 'tocb' if level == 0 else 'toc'
        left = .16*inch if level else 0
        ps = ParagraphStyle('tocleft', parent=ST[style], leftIndent=left)
        rows.append([Paragraph(escape(label), ps), Paragraph(str(page), ParagraphStyle('tocnum', parent=ST['toc'], alignment=TA_RIGHT))])
    toc_table = Table(rows, colWidths=[5.75*inch, .3*inch], hAlign='LEFT')
    toc_table.setStyle(TableStyle([
        ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),
        ('RIGHTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),0),
    ]))
    out.append(toc_table)
    return out


def make_page(page):
    out = [pdf_head(page['title'])]
    for sub, paragraphs in page['sections']:
        if sub: out.append(pdf_head(sub, 2))
        ref_mode = (page['title'] == 'KẾT LUẬN' and sub == 'DANH MỤC TÀI LIỆU THAM KHẢO')
        out.extend(para(text, 'ref' if ref_mode else 'body') for text in paragraphs)
    return out


def page_number(canvas, doc):
    if doc.page > 1:
        canvas.saveState()
        canvas.setFont('TNR', 11)
        canvas.setFillColor(colors.black)
        canvas.drawCentredString(A4[0] / 2, 1.05*cm, str(doc.page))
        canvas.restoreState()


def build_pdf():
    # First, verify each explicitly paginated frame can hold its contents.
    left, right, top, bottom = 3.5*cm, 2*cm, 2.5*cm, 2.5*cm
    frame = Frame(left, bottom, A4[0]-left-right, A4[1]-top-bottom,
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id='check')
    story_pages = [make_cover(), make_members_toc()] + [make_page(p) for p in PAGES]
    for index, flow in enumerate(story_pages, 1):
        from io import BytesIO
        from reportlab.pdfgen.canvas import Canvas
        probe = flow[:]
        f2 = Frame(left, bottom, A4[0]-left-right, A4[1]-top-bottom,
                   leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id='fit2')
        f2.addFromList(probe, Canvas(BytesIO(), pagesize=A4))
        used_pt = (A4[1] - top) - f2._y
        print(f'PAGE FIT {index}: {len(flow)-len(probe)}/{len(flow)} elements, {used_pt:.1f} pt used')
        if probe:
            raise RuntimeError(f'Page {index} overflowed by {len(probe)} flowable(s)')

    doc = SimpleDocTemplate(str(PDF_OUT), pagesize=A4,
        leftMargin=left, rightMargin=right, topMargin=top, bottomMargin=bottom,
        title='Kinh doanh dịch vụ sàn giao dịch bất động sản - Nhóm 5',
        author='Nhóm 5', subject='Báo cáo Khoa Luật - 10 trang nội dung')
    story = []
    for idx, page_flow in enumerate(story_pages):
        if idx: story.append(PageBreak())
        story.extend(page_flow)
    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)


if __name__ == '__main__':
    build_docx()
    build_pdf()
    from pypdf import PdfReader
    print('PDF_PAGE_COUNT', len(PdfReader(str(PDF_OUT)).pages))
    print('OUTPUT', FULL_DOCX, FULL_DOCX.stat().st_size)
    print('OUTPUT', PDF_OUT, PDF_OUT.stat().st_size)
