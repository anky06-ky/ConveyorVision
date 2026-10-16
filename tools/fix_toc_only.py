from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

SOURCE = Path(r'D:\Nhu\BT_03\Bài Tập nhóm 5-lớp 01.docx')
OUTPUT = Path(r'D:\Nhu\BT_03\Bai_tap_nhom_5_Chi_chinh_muc_luc.docx')
doc = Document(SOURCE)

paragraphs = doc.paragraphs
toc_title_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == 'MỤC LỤC')
body_start_index = next(i for i in range(toc_title_index + 1, len(paragraphs))
                        if paragraphs[i].text.strip() == 'PHẦN MỞ ĐẦU')

# Clear source outline markers and assign them only to actual headings. This
# changes navigation metadata only; the visible report text remains untouched.
outline_levels = {
    'DANH SÁCH THÀNH VIÊN': 0,
    'LỜI CẢM ƠN': 0,
    'LỜI CAM ĐOAN': 0,
    'PHẦN MỞ ĐẦU': 0,
    'I.Lý do chọn đề tài': 1,
    'Mục đích và nhiệm vụ nghiên cứu': 1,
    'Mục đích nghiên cứu': 2,
    'Nhiệm vụ nghiên cứu': 2,
    '2.1. Nhiệm vụ 1: Hệ thống hóa cơ sở lý luận và làm rõ các quy định pháp luật hiện hành về kinh doanh dịch vụ sàn giao dịch bất động sản': 2,
    '2.2. Nhiệm vụ 2: Đánh giá thực trạng áp dụng pháp luật, chỉ ra những hạn chế, bất cập và nguyên nhân': 2,
    '2.3. Nhiệm vụ 3: Đề xuất hệ thống giải pháp hoàn thiện pháp luật và nâng cao hiệu quả thi hành': 2,
    'Đối tượng và phạm vi nghiên cứu': 1,
    'Đối tượng nghiên cứu': 2,
    'Phạm vi nghiên cứu': 2,
    'Phương pháp nghiên cứu': 1,
    'Phương pháp phân tích pháp luật': 2,
    '2. Phương pháp tổng hợp pháp luật': 2,
    '3. Phương pháp so sánh pháp luật': 2,
    '4. Phương pháp nghiên cứu trường hợp (Case study)': 2,
    'PHẦN NỘI DUNG': 0,
    'Nội dung pháp luật hiện hành trong kinh doanh dịch vụ sàn giao dịch bất động sản': 1,
    'Cơ sở pháp lý': 2,
    'Khái niệm': 2,
    'Điều kiện thành lập và hoạt động sàn giao dịch bất động sản': 2,
    'Nội dung và phạm vi hoạt động': 2,
    '5. Quyền và nghĩa vụ của sàn giao dịch kinh doanh bất động sản': 2,
    'Những vấn đề thực tiễn liên quan đến kinh doanh dịch vụ sàn giao dịch bất động sản': 1,
    '1. Tình trạng công khai thông tin pháp lý bất động sản chưa đầy đủ, thiếu minh bạch trên sàn giao dịch': 2,
    'Hiện tượng môi giới thổi giá, tạo giao dịch ảo và thao túng tâm lý người mua': 2,
    '3. Tranh chấp về phí dịch vụ môi giới và phí giao dịch qua sàn': 2,
    '4. Trách nhiệm của sàn khi bất động sản không đủ điều kiện giao dịch': 2,
    '5. Hoạt động của sàn giao dịch bất động sản trực tuyến': 2,
    '6. Khó khăn trong quản lý, giám sát và xử lý vi phạm của các sàn': 2,
    '7. Khó khăn khi áp dụng quy định mới về kinh doanh dịch vụ sàn giao dịch bất động sản': 2,
    'Những vấn đề bất cập pháp luật phát sinh trong thực tiễn áp dụng pháp luật trong kinh doanh dịch vụ sàn giao dịch bất động sản': 1,
    '1. Khó khăn trong việc kiểm tra tính chính xác của thông tin bất động sản': 2,
    '2. Khó xác định trách nhiệm khi nhân viên môi giới cung cấp thông tin sai': 2,
    '3. Hoạt động giao dịch bất động sản trực tuyến phát sinh nhiều vấn đề mới': 2,
    '4. Công tác quản lý, kiểm tra hoạt động của sàn còn gặp khó khăn': 2,
    '5. Sự liên quan của nhiều quy định pháp luật gây khó khăn trong quá trình áp dụng': 2,
    'KẾT LUẬN': 0,
    'Kết luận': 1,
    'Đề xuất và hướng hoàn thiện': 1,
    'DANH MỤC TÀI LIỆU THAM KHẢO': 0,
}

for i, paragraph in enumerate(paragraphs):
    if i == toc_title_index:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.3
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.rFonts
            if rfonts is None:
                rfonts = OxmlElement('w:rFonts')
                rpr.insert(0, rfonts)
            for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                rfonts.set(qn(f'w:{key}'), 'Times New Roman')
    if i >= body_start_index:
        ppr = paragraph._p.get_or_add_pPr()
        for old in list(ppr.findall(qn('w:outlineLvl'))):
            ppr.remove(old)
        text = ' '.join(paragraph.text.split())
        if text in outline_levels:
            outline = OxmlElement('w:outlineLvl')
            outline.set(qn('w:val'), str(outline_levels[text]))
            ppr.append(outline)

# Define the generated TOC styles as black Times New Roman, 13 pt, 1.3 spacing.
for level in range(1, 4):
    name = f'TOC {level}'
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(13)
    style.font.bold = level == 1
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{key}'), 'Times New Roman')
    fmt = style.paragraph_format
    fmt.line_spacing = 1.3
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.left_indent = Cm((level - 1) * 0.5)
    fmt.first_line_indent = Cm(0)
    fmt.tab_stops.add_tab_stop(Inches(6.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

# Add a real Word TOC field. Word updates the page numbers on opening because
# the document setting below requests a field refresh.
field_p = doc.add_paragraph()
field_p.paragraph_format.line_spacing = 1.3
field_p.paragraph_format.space_after = Pt(0)
begin_run = OxmlElement('w:r')
begin = OxmlElement('w:fldChar')
begin.set(qn('w:fldCharType'), 'begin')
begin.set(qn('w:dirty'), 'true')
begin_run.append(begin)
instr_run = OxmlElement('w:r')
instr = OxmlElement('w:instrText')
instr.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
instr.text = ' TOC \\o "1-3" \\z \\u '
instr_run.append(instr)
separate_run = OxmlElement('w:r')
separate = OxmlElement('w:fldChar')
separate.set(qn('w:fldCharType'), 'separate')
separate_run.append(separate)
end_run = OxmlElement('w:r')
end = OxmlElement('w:fldChar')
end.set(qn('w:fldCharType'), 'end')
end_run.append(end)
for run in (begin_run, instr_run, separate_run, end_run):
    field_p._p.append(run)

# Place the field immediately after the TOC title and start the unchanged body
# on the following page.
title = paragraphs[toc_title_index]
title._p.addnext(field_p._p)
body_heading = paragraphs[body_start_index]
body_heading.paragraph_format.page_break_before = True

settings = doc.settings._element
update_fields = settings.find(qn('w:updateFields'))
if update_fields is None:
    update_fields = OxmlElement('w:updateFields')
    settings.append(update_fields)
update_fields.set(qn('w:val'), 'true')

doc.core_properties.title = 'Bài tập nhóm 5 - chỉ tinh chỉnh mục lục'
doc.save(OUTPUT)
print('Saved', OUTPUT)
print('Source paragraphs preserved:', len(paragraphs))
print('Heading outline levels assigned:', sum(1 for p in paragraphs[body_start_index:] if 'w:outlineLvl' in p._p.xml))
